from datetime import datetime
import csv
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference


# Converts a weight from pounds to kilograms.
def convertData(weight):
    return weight / 2.205


# Inserts comma-separated data into the specified CSV file.
def insertData(file_path, data):
    try:
        with open(file_path, "a") as file:
            file.write(data + "\n")
        return True
    except Exception as error:
        print("Error writing to file:", error)
        return False


# Displays the contents of the specified CSV file.
def viewData(file_path):
    try:
        print("The file", file_path)

        with open(file_path, "r") as file:
            print(file.read())
    except Exception as error:
        print("Error reading file:", error)


# Gets weight data from the user, converts it, and saves it to the CSV file.
def getInput():
    entries = int(input("How many entries are you inputting? "))

    for entry in range(entries):
        date = input("Enter a date: ")
        weight = float(input("Enter the weight in pounds for the inputted date: "))

        # Calls convertData with weight as the argument and returns the converted weight in kilograms.
        converted_weight = convertData(weight)

        data = str(date) + "," + str(weight) + "," + str(converted_weight)

        try:
            if insertData("ZooData.csv", data):
                print("The following data was saved at", datetime.now())
                print(data)
                print()
        except Exception as error:
            print("Error saving data:", error)


# Creates an Excel spreadsheet and chart from the CSV data.
# Arguments: file_path (string) is the CSV path and chart_type (string) is "bar" or "line".
# Returns: None.
def createChart(file_path, chart_type):
    print("Choose the data source:")
    print("1. Pounds")
    print("2. Kilograms")

    data_choice = input("Enter your selection: ")

    if data_choice not in ["1", "2"]:
        print("Error: Invalid data source selected.")
        return

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Weight Data"

    worksheet.append(["Date", "Pounds", "Kilograms"])

    try:
        with open(file_path, "r") as file:
            reader = csv.reader(file)

            for row in reader:
                if len(row) >= 3:
                    date = row[0]
                    pounds = float(row[1])
                    kilograms = float(row[2])

                    worksheet.append([date, pounds, kilograms])

    except Exception as error:
        print("Error reading file:", error)
        return

    if data_choice == "1":
        data_column = 2
        y_axis_title = "Weight (Pounds)"
    else:
        data_column = 3
        y_axis_title = "Weight (Kilograms)"

    if chart_type == "bar":
        chart = BarChart()
    else:
        chart = LineChart()

    values = Reference(
        worksheet,
        min_col=data_column,
        min_row=1,
        max_row=worksheet.max_row
    )

    dates = Reference(
        worksheet,
        min_col=1,
        min_row=2,
        max_row=worksheet.max_row
    )

    chart.add_data(values, titles_from_data=True)
    chart.set_categories(dates)

    chart.x_axis.title = "Date"
    chart.y_axis.title = y_axis_title
    chart.title = "Duspul5115" + datetime.now().strftime("%m/%d/%Y")

    worksheet.add_chart(chart, "E2")

    workbook.save("final.xlsx")

    print("Report successfully created as final.xlsx")


# Asks the user which chart type to create and calls createChart.
# Argument: file_path (string) is the path to the CSV data file.
# Returns: None.
def generateReport(file_path):
    print("Choose a chart type:")
    print("1. Bar Chart")
    print("2. Line Chart")

    chart_choice = input("Enter your selection: ")

    if chart_choice == "1":
        createChart(file_path, "bar")
    elif chart_choice == "2":
        createChart(file_path, "line")
    else:
        print("Error: Invalid chart type selected.")


print("Duspul5115 Spreadsheet Automation Menu")

menu_options = [
    "1. Input Data",
    "2. View Current Data",
    "3. Generate Report"
]

print("Choose a number from the following options")

for option in menu_options:
    print(option)

choice = input("Enter your selection: ")

if choice == "1":
    print("You selected", choice, "at", datetime.now())
    getInput()
elif choice == "2":
    print("You selected", choice, "at", datetime.now())
    viewData("ZooData.csv")
elif choice == "3":
    print("You selected", choice, "at", datetime.now())
    generateReport("ZooData.csv")
else:
    print("Error: The chosen functionality is not implemented yet")